"""
finlib.xlsx.styles — the house style for the workbook. Reusable across projects.

Implements the finance convention an experienced reviewer reads in half a second:
    BLUE font  = hard-coded input (a sourced or assumed number)
    BLACK font = formula (a computed cell)
    GREEN font = link to another sheet
Assumptions are additionally flagged so they can never be mistaken for sourced figures.
"""
from __future__ import annotations
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle

# --- palette ---
NAVY = "1F3A5F"
LIGHT = "DCE6F1"
GREY = "808080"
BLUE_INPUT = "0000FF"
GREEN_LINK = "1F7A1F"
AMBER = "9C6500"
AMBER_FILL = "FFF2CC"

MONEY_FMT = "#,##0;(#,##0)"        # thousands, negatives in parentheses
MONEY0 = "#,##0;(#,##0);\"-\""
PCT_FMT = "0.0%"
RATIO_FMT = "0.00"
DAYS_FMT = "0"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
TOP_BORDER = Border(top=Side(style="thin", color="000000"))
DBL_TOP = Border(top=Side(style="double", color="000000"))


def _font(**kw):
    kw.setdefault("size", 10)
    return Font(name="Calibri", **kw)


def title_font():   return _font(size=16, bold=True, color="FFFFFF")
def header_font():  return _font(bold=True, color="FFFFFF")
def label_font():   return _font(color="000000")
def input_font():   return _font(color=BLUE_INPUT)          # sourced/assumed hard number
def formula_font(): return _font(color="000000")            # computed
def link_font():    return _font(color=GREEN_LINK)          # cross-sheet reference
def subtotal_font(): return _font(bold=True, color="000000")
def note_font():    return _font(size=8, italic=True, color=GREY)
def assumption_font(): return _font(color=BLUE_INPUT, italic=True)

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor=LIGHT)
ASSUMPTION_FILL = PatternFill("solid", fgColor=AMBER_FILL)

LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
RIGHT = Alignment(horizontal="right", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def write(ws, cell, value, *, font=None, fmt=None, fill=None, align=None, border=None):
    c = ws[cell]
    c.value = value
    c.font = font or label_font()
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = align or LEFT
    if border:
        c.border = border
    return c


def header_row(ws, row, headers, start_col=1):
    from openpyxl.utils import get_column_letter
    for i, h in enumerate(headers):
        col = get_column_letter(start_col + i)
        write(ws, f"{col}{row}", h, font=header_font(), fill=HEADER_FILL,
              align=CENTER, border=BORDER)


def title_block(ws, title, subtitle=None):
    ws.merge_cells("A1:H1")
    write(ws, "A1", title, font=title_font(), fill=HEADER_FILL, align=LEFT)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells("A2:H2")
        write(ws, "A2", subtitle, font=note_font())
